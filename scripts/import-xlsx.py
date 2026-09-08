#!/usr/bin/env python3
"""Parse data/舆情.xlsx into src/lib/feedback.json."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "舆情.xlsx"
SUPPLEMENTS = sorted((ROOT / "data").glob("舆情补*.xlsx"))
OUT = ROOT / "src" / "lib" / "feedback.json"

RANGE_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})-(\d{1,2})\.(\d{1,2})$")
HEADER_MARKERS = ("反馈类型",)
YEAR = 2026


def parse_count(value) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).replace("条", "").replace(",", "").strip()
    try:
        return int(float(text))
    except ValueError:
        return 0


def parse_share(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        n = float(value)
        return n if n <= 1 else n / 100
    text = str(value).replace("%", "").strip()
    try:
        n = float(text)
        return n / 100 if n > 1 else n
    except ValueError:
        return 0.0


def clean_name(raw: str) -> str:
    text = str(raw).strip()
    match = re.search(r"（([^）]+)）", text)
    if match and any("\u4e00" <= ch <= "\u9fff" for ch in match.group(1)):
        return match.group(1).strip()
    return re.sub(r"（[^）]*）", "", text).strip() or text


def slug(name: str) -> str:
    keep = re.sub(r"[^\w\u4e00-\u9fff]+", "-", name).strip("-").lower()
    return keep or "item"


def to_range(label: str) -> tuple[str, str]:
    m = RANGE_RE.match(label)
    assert m
    m1, d1, m2, d2 = (int(x) for x in m.groups())
    long = f"{YEAR}.{m1:02d}.{d1:02d} — {m2:02d}.{d2:02d}"
    short = f"{m1:02d}.{d1:02d} – {m2:02d}.{d2:02d}"
    return long, short


def split_quotes(raw) -> list[str]:
    if not raw:
        return []
    text = str(raw)
    text = text.replace("<br>", "\n").replace("##Other Content##", "\n")
    text = re.sub(r"##[^#\n]*##", "\n", text)
    parts = re.split(r"\n+| \| |\s\|\s|(?<=[。！？.!?])\s*(?=\d+\.)", text)
    cleaned: list[str] = []
    for part in parts:
        item = part.strip().strip('"').strip("'").strip()
        item = re.sub(r"^\d+\.\s*", "", item)
        item = re.sub(r"（基于常见.*?）", "", item)
        item = re.sub(r"\s+", " ", item).strip(" |-")
        if len(item) < 12:
            continue
        if item in cleaned:
            continue
        cleaned.append(item)
        if len(cleaned) >= 3:
            break
    return cleaned


def is_header(cell) -> bool:
    if not isinstance(cell, str):
        return False
    return cell.startswith("反馈类型")


def parse_sheet(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    weeks: list[dict] = []
    current: dict | None = None

    def close():
        nonlocal current
        if current and current["categories"]:
            weeks.append(current)
        current = None

    for row in ws.iter_rows(values_only=True):
        a = row[0]
        if a is None:
            close()
            continue
        if isinstance(a, str) and RANGE_RE.match(a.strip()):
            close()
            label = a.strip()
            long_range, short = to_range(label)
            current = {
                "id": f"w-{label}",
                "label": label,
                "range": long_range,
                "rangeShort": short,
                "categories": [],
                "quotes": [],
            }
            continue
        if current is None or is_header(a):
            continue

        name = clean_name(str(a))
        count = parse_count(row[1])
        share = parse_share(row[2])
        summary = str(row[3] or "").strip()
        suggestion = str(row[5] or "").strip() if len(row) > 5 and row[5] else ""
        cat_id = slug(name)
        current["categories"].append(
            {
                "id": cat_id,
                "name": name,
                "count": count,
                "share": round(share, 4),
                "summary": summary,
                "suggestion": suggestion,
            }
        )
        for i, quote in enumerate(split_quotes(row[4]), 1):
            current["quotes"].append(
                {
                    "id": f"{current['id']}-{cat_id}-{i}",
                    "categoryId": cat_id,
                    "category": name,
                    "text": quote,
                }
            )

    close()
    for week in weeks:
        total = sum(c["count"] for c in week["categories"])
        week["total"] = total
        top = week["categories"][0]
        second = week["categories"][1] if len(week["categories"]) > 1 else None
        share_pct = round(top["share"] * 100, 1)
        week["narrative"] = (
            f"本周有效反馈 {total} 条。占比最高的是「{top['name']}」{top['count']} 条"
            f"（{share_pct}%）"
            + (
                f"，其次是「{second['name']}」{second['count']} 条。"
                if second
                else "。"
            )
        )
        source = top["suggestion"] or top["summary"]
        week["judgment"] = re.split(r"[。；\n]", source)[0].strip()[:120]
        if week["judgment"] and not week["judgment"].endswith("。"):
            week["judgment"] += "。"
    weeks.reverse()
    return weeks


def week_sort_key(week: dict) -> tuple[int, int]:
    m = RANGE_RE.match(week["label"])
    if not m:
        return (0, 0)
    return (int(m.group(1)), int(m.group(2)))


def main() -> None:
    by_id: dict[str, dict] = {}
    for path in [XLSX, *SUPPLEMENTS]:
        if not path.exists():
            continue
        for week in parse_sheet(path):
            by_id[week["id"]] = week
            print(f"loaded {week['label']} ({week['total']}) from {path.name}")

    weeks = sorted(by_id.values(), key=week_sort_key, reverse=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps({"year": YEAR, "weeks": weeks}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"wrote {len(weeks)} weeks -> {OUT}")
    print("latest", weeks[0]["label"], weeks[0]["total"])
    print("oldest", weeks[-1]["label"], weeks[-1]["total"])
    print("totals", sum(w["total"] for w in weeks))


if __name__ == "__main__":
    main()
